#!/usr/bin/env python3
"""Build data/vulnerability_curves/ from the VulnerabilityCurves source repo.

The source (https://github.com/SebKrantz/VulnerabilityCurves) exports one CSV
per curve from the supplementary tables of

    Nirandjan, S. et al. (2024) A spatially-explicit harmonized global dataset
    of critical infrastructure vulnerability curves. NHESS 24, 4341-4369.
    https://doi.org/10.5194/nhess-24-4341-2024

This script turns those 394 files into the library the app ships:

  * **Families are merged.** 88 curves come as a central file plus `a` (lower)
    and `b` (upper) bound files on identical intensity grids. Those become one
    four-column CSV `intensity,lower,central,upper` — the layout
    `parse_vulnerability_curve_data` already reads, which means the app draws
    error bars and reports a damage-cost band without any further work.
  * **Metadata is joined in** from Table D1 (asset description, characteristics,
    geography, source, derivation method) and written to `curves_index.csv`,
    the file the assistant's curve_library module searches.
  * **Hazards that match an app layer's units are promoted**; the rest go under
    `reference/`, where they are searchable but never auto-selected.
  * **Table D3 costs** become `asset_costs.csv`, the replacement-value reference.

Re-run after updating the source repo:

    python scripts/build_curve_library.py [--source PATH] [--check]

`--check` rebuilds into a temp directory and diffs against the committed
library instead of overwriting it.
"""

from __future__ import annotations

import argparse
import csv
import re
import shutil
import sys
from collections import defaultdict
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

REPO = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = REPO.parent / "VulnerabilityCurves"
TARGET = REPO / "data" / "vulnerability_curves"

# Source hazard folder -> (library folder, app hazard family, matching app layers).
# Only these three share an intensity metric AND a unit with a layer in
# data/hazard_layers.csv, so only these can be applied without conversion.
PROMOTED = {
    "flood": ("flood_mm", "flood", "Flood Hazard 25/50/100 Years (all climate variants)"),
    "earthquake_pga_cms2": (
        "earthquake_pga_cms2",
        "earthquake_pga",
        "Peak Ground Acceleration PGA - 250/475/975 Years",
    ),
    "wind_kmh": (
        "cyclone_wind_kmh",
        "cyclone_wind",
        "Tropical Cyclone Wind 25/50/100 Years (with and without climate change)",
    ),
}

# Intensity anchors reported in the index, in each family's own unit. They let
# the assistant compare candidate curves without opening any file.
ANCHORS = {
    "flood": [100, 500, 1000, 2000],
    "earthquake_pga": [100, 200, 500, 1000],
    "cyclone_wind": [120, 180, 250],
}

INTENSITY_METRIC = {
    "flood_mm": ("inundation depth", "mm"),
    "earthquake_pga_cms2": ("peak ground acceleration", "cm/s2"),
    "cyclone_wind_kmh": ("10 m gust wind speed", "km/h"),
    "earthquake_mmi": ("modified Mercalli intensity", "MMI"),
    "earthquake_pgv_cms": ("peak ground velocity", "cm/s"),
    "earthquake_mgv_cms": ("maximum ground velocity", "cm/s"),
    "earthquake_sa_g": ("spectral acceleration", "g"),
    "earthquake_sd_cm": ("spectral displacement", "cm"),
    "earthquake_pgv2pga_cm": ("PGV^2/PGA", "cm"),
    "landslide_pgd_cm": ("permanent ground deformation", "cm"),
    "landslide_precipitation_mm": ("triggering precipitation", "mm"),
    "landslide_area_m2": ("landslide area", "m2"),
    "landslide_slope_velocity_mmyr": ("slope velocity", "mm/year"),
}

# Curves whose output is not a damage factor in [0, 1]. They stay in the
# library for reference but must never be fed to the app's damage calculation.
NOT_DAMAGE_FACTOR = {
    "earthquake_mmi": "repair rate, not a damage factor",
    "earthquake_mgv_cms": "damaged pipes per km, not a damage factor",
    "earthquake_pgv2pga_cm": "repair rate, not a damage factor",
}

D1_SHEETS = [
    "Energy",
    "Transportation",
    "Telecommunication",
    "Water",
    "Waste",
    "Education & Health",
]


def _cell(value) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ").strip()
    return "" if text.lower() in ("nan", "none", "n/a") else text


def load_d1(source: Path) -> dict[str, dict]:
    """curve ID -> metadata row from Table D1."""
    path = source / "curves_excel" / "Table_D1_Summary_CI_Vulnerability_Data_V1.0.0.xlsx"
    book = openpyxl.load_workbook(path, read_only=True)
    out: dict[str, dict] = {}
    for sheet in D1_SHEETS:
        rows = list(book[sheet].iter_rows(values_only=True))
        header = [_cell(c) for c in rows[1]]
        for row in rows[2:]:
            if not row:
                continue
            rec = dict(zip(header, [_cell(c) for c in row]))
            rec["_reference"] = _cell(row[0])
            curve_id = rec.get("ID number", "")
            if curve_id:
                out.setdefault(curve_id, rec)
    book.close()
    return out


def load_d3(source: Path) -> list[dict]:
    """Table D3 replacement/maximum-damage costs, one row per published figure."""
    path = source / "curves_excel" / "Table_D3_Costs_V1.0.0.xlsx"
    book = openpyxl.load_workbook(path, read_only=True)
    rows = list(book["Cost_Database"].iter_rows(values_only=True))
    book.close()
    header = [_cell(c) for c in rows[0]]
    out = []
    for row in rows[1:]:
        rec = dict(zip(header, [_cell(c) for c in row]))
        if not rec.get("Infrastructure description"):
            continue
        out.append(rec)
    return out


def parse_family_token(stem: str) -> tuple[str, str] | None:
    """('e11', 'a') from a source filename stem; bound is '' for the central."""
    parts = stem.split("_")
    token = parts[-3] if parts[-1] == "boundary" else parts[-1]
    # One file breaks the convention: ..._e141_a.csv rather than ..._e141a.csv.
    if token in ("a", "b") and len(parts) >= 2:
        token = parts[-2] + token
    match = re.fullmatch(r"([a-z]+\d+)([ab])?", token)
    if not match:
        return None
    return match.group(1), match.group(2) or ""


# Curve IDs come from Table D1, whose keys are unambiguous ('E21.18', not the
# 'E2.118' a naive digit split would produce). One curve is absent from D1
# altogether; its ID and description are read off the Table D2 sheet header.
ID_OVERRIDES = {
    "e2125": {
        "ID number": "E21.25",
        "Infrastructure description": "Transmission and distribution pipelines",
        "Additional characteristics": "Buried pipelines",
        "_reference": "FEMA (2020) Hazus Earthquake Model",
        "Geographical application": "USA",
    },
}


def read_points(path: Path) -> tuple[np.ndarray, np.ndarray]:
    frame = pd.read_csv(path)
    return frame.iloc[:, 0].to_numpy(float), frame.iloc[:, 1].to_numpy(float)


def anchors_for(family: str, intensity: np.ndarray, central: np.ndarray) -> str:
    """'100:0.02; 500:0.21; 2000:0.55~' — '~' marks a clamped extrapolation."""
    wanted = ANCHORS.get(family)
    if not wanted:
        lo, hi = float(intensity.min()), float(intensity.max())
        wanted = [round(lo + f * (hi - lo), 3) for f in (0.25, 0.5, 0.75, 1.0)]
    parts = []
    for x in wanted:
        y = float(np.interp(x, intensity, central))
        clamped = "~" if x > intensity.max() or x < intensity.min() else ""
        parts.append(f"{x:g}:{y:.4g}{clamped}")
    return "; ".join(parts)


def build(source: Path, target: Path) -> dict:
    d1 = load_d1(source)
    d1_by_token = {re.sub(r"[^a-z0-9]", "", k.lower()): v for k, v in d1.items()}

    families: dict[tuple[str, str, str], dict[str, Path]] = defaultdict(dict)
    for path in sorted((source / "curves_csv").rglob("*.csv")):
        parsed = parse_family_token(path.stem)
        if parsed is None:
            raise SystemExit(f"unparseable curve filename: {path}")
        token, bound = parsed
        families[(path.parent.parent.name, path.parent.name, token)][bound or "c"] = path

    # Clear the generated content but keep README.md, which is written by hand.
    if target.exists():
        for child in target.iterdir():
            if child.name == "README.md":
                continue
            shutil.rmtree(child) if child.is_dir() else child.unlink()
    target.mkdir(parents=True, exist_ok=True)

    index: list[dict] = []
    unmatched: list[str] = []

    for (source_hazard, sector, token), members in sorted(families.items()):
        promoted = PROMOTED.get(source_hazard)
        folder = promoted[0] if promoted else source_hazard
        family = promoted[1] if promoted else ""
        metric, unit = INTENSITY_METRIC[folder]

        # Metadata: exact ID first, then the base ID for an a/b bound variant.
        base_token = re.sub(r"[ab]$", "", token)
        meta = (
            d1_by_token.get(token)
            or d1_by_token.get(base_token)
            or ID_OVERRIDES.get(base_token)
            or {}
        )
        curve_id = meta.get("ID number", "")
        if not curve_id:
            unmatched.append(token)
            curve_id = token.upper()

        central_path = members.get("c")
        lower_path, upper_path = members.get("a"), members.get("b")
        has_bounds = lower_path is not None and upper_path is not None

        if central_path is not None:
            intensity, central = read_points(central_path)
            central_basis = "published"
            name_source = central_path
        else:
            # Three families publish only the bounds; the midpoint is the
            # honest central estimate and the index records that it is derived.
            intensity, low = read_points(lower_path)
            _, high = read_points(upper_path)
            central = (low + high) / 2.0
            central_basis = "midpoint of published bounds"
            name_source = lower_path

        columns = {"intensity": intensity, "central": central}
        bounds_note = ""
        if has_bounds:
            lo_x, lo_y = read_points(lower_path)
            hi_x, hi_y = read_points(upper_path)
            if not (np.allclose(lo_x, intensity) and np.allclose(hi_x, intensity)):
                raise SystemExit(f"bound grids differ from the central grid for {token}")
            # Nine families publish bounds that do not bracket their own central
            # curve. Below 0.001 that is float noise and is clamped away; above
            # it the band contradicts the central estimate, and shipping it would
            # hand the app a damage range whose "upper" is lower than its middle.
            # Those bounds are dropped and the reason recorded.
            violation = max(float((lo_y - central).max()), float((central - hi_y).max()), 0.0)
            if violation <= 1e-3:
                columns["lower"] = np.minimum(lo_y, central)
                columns["upper"] = np.maximum(hi_y, central)
            else:
                has_bounds = False
                bounds_note = (
                    "published bounds discarded: they do not bracket the published "
                    f"central curve (by up to {violation:.3f})"
                )

        # Filename: the source's descriptive stem with the bound suffix dropped.
        stem = re.sub(r"_[a-z]+\d+[ab]?(_(lower|upper)_boundary)?$", "", name_source.stem)
        out_dir = target / ("reference" if not promoted else "") / folder / sector
        out_dir.mkdir(parents=True, exist_ok=True)
        out_path = out_dir / f"{stem}_{token}.csv"

        with out_path.open("w", newline="") as handle:
            writer = csv.writer(handle, lineterminator="\n")
            if has_bounds:
                writer.writerow(["intensity", "lower", "central", "upper"])
                rows = zip(intensity, columns["lower"], central, columns["upper"])
            else:
                writer.writerow(["intensity", "proportion_destroyed"])
                rows = zip(intensity, central)
            for row in rows:
                writer.writerow([f"{v:.6g}" for v in row])

        index.append(
            {
                "curve_id": curve_id,
                "hazard": family or folder,
                "app_layer_match": promoted[2] if promoted else "",
                "intensity_metric": metric,
                "unit": unit,
                "sector": sector,
                "asset": meta.get("Infrastructure description", ""),
                "characteristics": meta.get("Additional characteristics", ""),
                "geography": meta.get("Geographical application", ""),
                "source": meta.get("_reference", ""),
                "derivation": meta.get("Derivation methodology", "")
                or meta.get("Derivation method", ""),
                "output": NOT_DAMAGE_FACTOR.get(folder, "damage factor 0-1"),
                "points": len(intensity),
                "intensity_min": f"{intensity.min():.6g}",
                "intensity_max": f"{intensity.max():.6g}",
                "damage_max": f"{central.max():.4g}",
                "has_bounds": "yes" if has_bounds else "no",
                "central_basis": central_basis,
                "bounds_note": bounds_note,
                "damage_at": anchors_for(family, intensity, central),
                "monotonic": "yes" if np.all(np.diff(central) >= -1e-9) else "no",
                "file": str(out_path.relative_to(target)),
            }
        )

    index.sort(key=lambda r: (r["hazard"], r["sector"], r["curve_id"]))
    with (target / "curves_index.csv").open("w", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(index[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(index)

    costs = load_d3(source)
    with (target / "asset_costs.csv").open("w", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(
            [
                "asset",
                "characteristics",
                "amount",
                "lower",
                "upper",
                "unit",
                "geography",
                "cost_feature",
                "source",
                "curve_ids",
            ]
        )
        for rec in costs:
            writer.writerow(
                [
                    rec.get("Infrastructure description", ""),
                    rec.get("Additional characteristics", ""),
                    rec.get("Amount", ""),
                    rec.get("Lower range", ""),
                    rec.get("Upper range", ""),
                    rec.get("Unit", ""),
                    rec.get("Geographical application", ""),
                    rec.get("Cost feature", ""),
                    rec.get("Source", ""),
                    rec.get("ID number", ""),
                ]
            )

    return {
        "families": len(families),
        "written": len(index),
        "parse_errors": validate(target, index),
        "bounded": sum(1 for r in index if r["has_bounds"] == "yes"),
        "bounds_discarded": [r["curve_id"] for r in index if r["bounds_note"]],
        "promoted": sum(1 for r in index if r["app_layer_match"]),
        "non_monotonic": [r["curve_id"] for r in index if r["monotonic"] == "no"],
        "unmatched_metadata": unmatched,
        "costs": len(costs),
    }


def validate(target: Path, index: list[dict]) -> list[str]:
    """Parse every emitted curve through the app's own reader.

    The library is only useful if `parse_vulnerability_curve_data` reads it, so
    that is what the build checks — not a reimplementation of it. Curves whose
    output is a repair rate rather than a damage factor are range-exempt.
    """
    sys.path.insert(0, str(REPO / "backend"))
    try:
        from app.utils.geospatial import parse_vulnerability_curve_data
    except ImportError as exc:  # backend deps absent; the files are still written
        return [f"skipped: {exc}"]

    problems: list[str] = []
    for row in index:
        path = target / row["file"]
        try:
            _, intensity, central, lower, upper = parse_vulnerability_curve_data(str(path))
        except Exception as exc:  # noqa: BLE001 — report, do not abort the build
            problems.append(f"{row['curve_id']}: unreadable ({exc})")
            continue
        if len(intensity) != int(row["points"]):
            problems.append(f"{row['curve_id']}: read {len(intensity)} points, index says {row['points']}")
        if np.any(np.diff(intensity) <= 0):
            problems.append(f"{row['curve_id']}: intensity not strictly increasing")
        if row["output"] == "damage factor 0-1" and (central.min() < 0 or central.max() > 1):
            problems.append(f"{row['curve_id']}: damage outside [0, 1]")
        if (lower is None) != (upper is None):
            problems.append(f"{row['curve_id']}: only one bound was parsed")
        if lower is not None:
            lo, hi = lower(intensity), upper(intensity)
            if np.any(lo > central + 1e-9) or np.any(hi < central - 1e-9):
                problems.append(f"{row['curve_id']}: bounds do not bracket the central curve")
    return problems


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument(
        "--check",
        action="store_true",
        help="build into a temp dir and diff against the committed library",
    )
    args = parser.parse_args()

    if not (args.source / "curves_csv").is_dir():
        print(f"source repo not found at {args.source}", file=sys.stderr)
        return 1

    if args.check:
        import tempfile

        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "vulnerability_curves"
            stats = build(args.source, out)
            import subprocess

            # README.md is hand-written and deliberately not regenerated;
            # .DS_Store is macOS litter.
            diff = subprocess.run(
                ["diff", "-r", "-x", "README.md", "-x", ".DS_Store", str(TARGET), str(out)],
                capture_output=True,
                text=True,
            )
            print(stats)
            if diff.returncode:
                print(diff.stdout[:4000])
                return 1
            print("library matches the source repo")
            return 0

    stats = build(args.source, TARGET)
    for key, value in stats.items():
        print(f"{key}: {value}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
